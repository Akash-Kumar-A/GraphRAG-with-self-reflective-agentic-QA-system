#!/usr/bin/env python3
"""
PSG public-site RAG crawler v2.

Fixes:
- Distinguishes explicit robots disallow from robots.txt fetch/parse errors.
- Uses a configurable, truthful User-Agent (do not impersonate a browser).
- Filters template URLs such as {{basicConfig.helpLink}}, mailto:, javascript:, and
  common broken generated links.
- Optional Playwright rendering for JavaScript-driven pages, ONLY after robots allows URL.
- Captures HTTP status, content type, extraction method and crawl status.
- Keeps all URLs and source metadata attached to extracted text.

Respect each site's robots.txt and terms. This tool never overrides a robots disallow
and does not bypass logins, CAPTCHAs, or other access controls.
"""
from __future__ import annotations
import argparse, hashlib, json, re, time, urllib.parse, urllib.robotparser
from collections import deque
from datetime import datetime, timezone
from pathlib import Path
import requests
from bs4 import BeautifulSoup

DEFAULT_STARTS = [
 "https://www.psgtech.edu/",
 "https://library.psgtech.ac.in/",
 "https://sports.psgtech.ac.in/",
 "https://alumni.psgtech.ac.in/",
 "https://iprcell.psgtech.ac.in/",
 "https://ir.psgtech.ac.in/",
 "https://journals.psgtech.ac.in/",
 "https://academicschedule.psgtech.ac.in/",
 "https://enquiry.psgtech.ac.in/",
 "https://iipc.psgtech.ac.in/",
 "https://cnce.psgtech.ac.in/",
 "https://csrc.psgtech.ac.in/",
 "https://aiconsortium.psgtech.ac.in/",
]
DEFAULT_DOMAINS = ["psgtech.edu", "psgtech.ac.in"]
SKIP_EXT = {".jpg",".jpeg",".png",".gif",".webp",".svg",".ico",".mp3",".mp4",
 ".mov",".avi",".zip",".rar",".7z",".exe",".dmg",".css",".js",".woff",".woff2",
 ".ttf",".eot"}
DOC_EXT = {".pdf",".docx",".pptx",".xlsx",".xls",".doc",".ppt"}
SKIP_PARTS = ["/login", "/logout", "/signin", "/sign-in", "/auth", "/account",
 "/cart", "/checkout", "/wp-admin", "/admin"]

def canonicalize(url):
    try:
        p=urllib.parse.urlsplit(url.strip())
        if p.scheme.lower() not in ("http","https") or not p.netloc: return None
        if any(x in url for x in ("{{","}}","${","<",">")): return None
        if p.username or p.password: return None
        path=re.sub(r"/{2,}","/",p.path or "/")
        # Remove common tracking params; keep other query params for genuine pages.
        q=urllib.parse.parse_qsl(p.query, keep_blank_values=True)
        q=[(k,v) for k,v in q if not k.lower().startswith("utm_") and k.lower() not in {"fbclid","gclid"}]
        query=urllib.parse.urlencode(q, doseq=True)
        return urllib.parse.urlunsplit((p.scheme.lower(),p.netloc.lower(),path,query,""))
    except Exception: return None

def host_allowed(url, domains):
    host=(urllib.parse.urlsplit(url).hostname or "").lower()
    return any(host==d or host.endswith("."+d) for d in domains)

def path_allowed(url):
    p=urllib.parse.urlsplit(url)
    low=(p.path+"?"+p.query).lower()
    if any(x in low for x in SKIP_PARTS): return False
    if Path(p.path.lower()).suffix in SKIP_EXT: return False
    if len(p.query)>300: return False
    return True

def robots_check(url, session, cache, ua):
    """Return (allowed, state, robots_url). Fail closed on robots fetch errors."""
    p=urllib.parse.urlsplit(url); origin=f"{p.scheme}://{p.netloc}"
    robots_url=origin+"/robots.txt"
    if origin not in cache:
        rp=urllib.robotparser.RobotFileParser()
        try:
            r=session.get(robots_url, timeout=15, allow_redirects=True)
            if r.status_code == 200:
                rp.set_url(robots_url); rp.parse(r.text.splitlines())
                state="robots_loaded"
            elif r.status_code in (401,403):
                rp.parse(["User-agent: *","Disallow: /"]); state=f"robots_http_{r.status_code}_fail_closed"
            elif r.status_code == 404:
                rp.parse(["User-agent: *","Allow: /"]); state="robots_404_no_rules"
            else:
                rp.parse(["User-agent: *","Disallow: /"]); state=f"robots_http_{r.status_code}_fail_closed"
        except requests.RequestException as e:
            rp.parse(["User-agent: *","Disallow: /"]); state="robots_fetch_error_fail_closed:"+str(e)
        cache[origin]=(rp,state,robots_url)
    rp,state,robots_url=cache[origin]
    try:
        ok=rp.can_fetch(ua,url)
    except Exception:
        ok=False
    return ok, state, robots_url

def html_extract(html, url):
    soup=BeautifulSoup(html,"html.parser")
    title=soup.title.get_text(" ",strip=True) if soup.title else ""
    for tag in soup(["script","style","noscript","svg","iframe"]): tag.decompose()
    # Keep nav/footer for broad archival; remove hidden UI labels when marked aria-hidden.
    for tag in soup.select('[aria-hidden="true"]'): tag.decompose()
    root=soup.find("main") or soup.find("article") or soup.find(id=re.compile("content|main",re.I)) or soup.body or soup
    text=root.get_text("\n",strip=True)
    text=re.sub(r"[ \t]+\n","\n",text); text=re.sub(r"\n[ \t]+","\n",text)
    text=re.sub(r"\n{3,}","\n\n",text).strip()
    links=[]
    for a in soup.find_all("a",href=True):
        href=a.get("href","").strip()
        if not href or href.startswith(("#","mailto:","tel:","javascript:","data:")): continue
        u=canonicalize(urllib.parse.urljoin(url,href))
        if u: links.append(u)
    return title,text,links

def extract_doc(data, ext):
    import io
    if ext==".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        return "\n\n".join(f"[PDF page {i}]\n{p.extract_text() or ''}" for i,p in enumerate(reader.pages,1))
    if ext==".docx":
        from docx import Document
        d=Document(io.BytesIO(data)); out=[p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows: out.append(" | ".join(c.text.strip().replace("\n"," / ") for c in row.cells))
        return "\n".join(out)
    if ext==".pptx":
        from pptx import Presentation
        prs=Presentation(io.BytesIO(data)); out=[]
        for i,s in enumerate(prs.slides,1):
            out.append(f"[Slide {i}]")
            for sh in s.shapes:
                if getattr(sh,"has_text_frame",False) and sh.text.strip(): out.append(sh.text.strip())
                if getattr(sh,"has_table",False):
                    for row in sh.table.rows: out.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(out)
    if ext in (".xlsx",".xls"):
        import pandas as pd
        book=pd.ExcelFile(io.BytesIO(data)); out=[]
        for sheet in book.sheet_names:
            df=pd.read_excel(book,sheet_name=sheet,header=None).fillna("")
            out.append(f"[Sheet: {sheet}]")
            out.extend(" | ".join(str(v) for v in row) for row in df.values.tolist())
        return "\n".join(out)
    return "[Legacy Office file: convert to a modern format for text extraction.]"

def render_playwright(url, timeout):
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None, "playwright_not_installed"
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True)
            page=browser.new_page()
            page.goto(url, wait_until="networkidle", timeout=timeout*1000)
            page.wait_for_timeout(1000)
            html=page.content()
            browser.close()
            return html, None
    except Exception as e:
        return None, "playwright_error:"+repr(e)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--out",default="psg_corpus")
    ap.add_argument("--max-pages",type=int,default=5000)
    ap.add_argument("--delay",type=float,default=1.5)
    ap.add_argument("--timeout",type=int,default=25)
    ap.add_argument("--domains",default=",".join(DEFAULT_DOMAINS))
    ap.add_argument("--start-urls",default=",".join(DEFAULT_STARTS))
    ap.add_argument("--user-agent",default="PSGPublicRAGCorpusBot/2.0 (public information indexing; contact site administrator)")
    ap.add_argument("--render-js",action="store_true",help="Use Playwright fallback on robots-allowed HTML pages with empty/short text")
    args=ap.parse_args()
    out=Path(args.out); out.mkdir(parents=True,exist_ok=True)
    snapshots=out/"pages"; snapshots.mkdir(exist_ok=True)
    domains=[d.strip().lower() for d in args.domains.split(",") if d.strip()]
    starts=[canonicalize(u) for u in args.start_urls.split(",") if u.strip()]
    starts=[u for u in starts if u]
    s=requests.Session(); s.headers.update({"User-Agent":args.user_agent})
    robots_cache={}; q=deque(starts); seen=set(); records=[]; failures=[]
    started=datetime.now(timezone.utc).isoformat()
    while q and len(seen)<args.max_pages:
        url=q.popleft()
        if not url or url in seen or not host_allowed(url,domains) or not path_allowed(url): continue
        seen.add(url)
        ok,robots_state,robots_url=robots_check(url,s,robots_cache,args.user_agent)
        if not ok:
            failures.append({"url":url,"reason":"blocked_by_robots_txt_or_robots_check","robots_state":robots_state,"robots_url":robots_url})
            continue
        try:
            r=s.get(url,timeout=args.timeout,allow_redirects=True)
            final=canonicalize(r.url) or url
            if r.status_code>=400:
                failures.append({"url":final,"status":r.status_code,"reason":"http_error"})
                time.sleep(max(.2,args.delay)); continue
            ctype=r.headers.get("Content-Type","").lower()
            ext=Path(urllib.parse.urlsplit(final).path.lower()).suffix
            title=""; body=""; links=[]; kind="html"; method="requests_html"
            if "text/html" in ctype or (not ctype and ext not in DOC_EXT):
                title,body,links=html_extract(r.text,final)
                if args.render_js and len(body)<100:
                    rendered,err=render_playwright(final,args.timeout)
                    if rendered:
                        t2,b2,l2=html_extract(rendered,final)
                        if len(b2)>len(body): title,body,links=t2,b2,l2; method="playwright_rendered_html"
                    elif err:
                        failures.append({"url":final,"reason":"js_render_fallback_unavailable","detail":err})
                if len(body)<30:
                    failures.append({"url":final,"reason":"no_or_very_little_text_extracted","extraction_method":method,"content_type":ctype})
            elif "application/pdf" in ctype or ext==".pdf" or ext in DOC_EXT:
                kind=ext.lstrip(".") or "document"; title=Path(urllib.parse.urlsplit(final).path).name
                try: body=extract_doc(r.content,ext)
                except Exception as e: failures.append({"url":final,"reason":"document_extraction_error","detail":repr(e)})
            else:
                failures.append({"url":final,"reason":"unsupported_content_type","content_type":ctype})
                time.sleep(max(.2,args.delay)); continue
            if body.strip():
                crawled=datetime.now(timezone.utc).isoformat()
                rec={"id":hashlib.sha1(final.encode()).hexdigest(),"url":final,"title":title or final,
                     "content_type":ctype,"kind":kind,"text":body.strip(),
                     "sha256_text":hashlib.sha256(body.encode()).hexdigest(),"crawled_at":crawled,
                     "source_domain":urllib.parse.urlsplit(final).hostname,"http_status":r.status_code,
                     "extraction_method":method}
                records.append(rec)
                (snapshots/(rec["id"]+".md")).write_text(
                    f"# {rec['title']}\n\nSource: {final}\nCrawled: {crawled}\nExtraction: {method}\n\n{body.strip()}\n",encoding="utf-8")
            for link in links:
                if host_allowed(link,domains) and path_allowed(link) and link not in seen: q.append(link)
        except Exception as e:
            failures.append({"url":url,"reason":"request_or_parse_error","detail":repr(e)})
        time.sleep(max(.2,args.delay))
    with (out/"psg_corpus.jsonl").open("w",encoding="utf-8") as f:
        for rec in records: f.write(json.dumps(rec,ensure_ascii=False)+"\n")
    (out/"crawl_failures.json").write_text(json.dumps(failures,ensure_ascii=False,indent=2),encoding="utf-8")
    manifest={"started_at":started,"finished_at":datetime.now(timezone.utc).isoformat(),
      "start_urls":starts,"allowed_domains":domains,"pages_and_documents_extracted":len(records),
      "unique_urls_attempted":len(seen),"failures_or_skips":len(failures),"output_jsonl":str(out/"psg_corpus.jsonl"),
      "robots_policy":"Disallowed URLs are not fetched; robots fetch errors fail closed.",
      "js_render_enabled":args.render_js,
      "note":"Check crawl_failures.json. A successful extraction is not proof of complete page coverage; scanned PDFs may require OCR."}
    (out/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(manifest,indent=2))

if __name__=="__main__": main()
